import io
from typing import Tuple
import fitz  # PyMuPDF
import docx

from backend.config import settings
from backend.utils.logger import logger


class ResumeParserError(Exception):
    """Custom exception raised when resume file parsing fails."""
    pass


class ResumeParser:
    """Service for local text extraction from PDF and DOCX resume files."""

    ALLOWED_EXTENSIONS = {".pdf", ".docx"}
    MAX_FILE_SIZE_BYTES = settings.MAX_RESUME_SIZE_MB * 1024 * 1024

    @classmethod
    def validate_file(cls, filename: str, content: bytes) -> str:
        """Validates file extension and size constraints. Returns extracted lower-case extension."""
        if not filename or "." not in filename:
            raise ResumeParserError("Invalid filename. File must have an extension (.pdf or .docx).")

        ext = "." + filename.rsplit(".", 1)[-1].lower()
        if ext not in cls.ALLOWED_EXTENSIONS:
            raise ResumeParserError(f"Unsupported file format '{ext}'. Allowed formats: PDF, DOCX.")

        if not content or len(content) == 0:
            raise ResumeParserError("Uploaded resume file is empty.")

        if len(content) > cls.MAX_FILE_SIZE_BYTES:
            raise ResumeParserError(
                f"File size exceeds maximum limit of {settings.MAX_RESUME_SIZE_MB}MB."
            )

        return ext

    @classmethod
    def parse_bytes(cls, filename: str, content: bytes) -> Tuple[str, str]:
        """
        Parses resume file content bytes and extracts raw text.
        Returns tuple of (extracted_text, file_extension).
        """
        ext = cls.validate_file(filename, content)

        try:
            if ext == ".pdf":
                raw_text = cls._extract_text_from_pdf(content)
            elif ext == ".docx":
                raw_text = cls._extract_text_from_docx(content)
            else:
                raise ResumeParserError(f"Unsupported format: {ext}")

            raw_text = raw_text.strip()
            if not raw_text or len(raw_text) < 20:
                raise ResumeParserError(
                    "Extracted text is empty or too short. The document may be scanned, image-only, or corrupted."
                )

            logger.info(f"Successfully extracted {len(raw_text)} characters from {filename}")
            return raw_text, ext
        except ResumeParserError:
            raise
        except Exception as e:
            logger.error(f"Unexpected error parsing file '{filename}': {e}")
            raise ResumeParserError(f"Failed to parse resume document: {str(e)}")

    @classmethod
    def _extract_text_from_pdf(cls, content: bytes) -> str:
        """Extracts text from PDF bytes using PyMuPDF (fitz)."""
        text_parts = []
        try:
            with fitz.open(stream=content, filetype="pdf") as doc:
                if doc.page_count == 0:
                    raise ResumeParserError("PDF document contains no pages.")
                for page in doc:
                    text_parts.append(page.get_text())
            return "\n".join(text_parts)
        except fitz.FileDataError:
            raise ResumeParserError("Corrupted or invalid PDF file.")
        except Exception as e:
            raise ResumeParserError(f"Error reading PDF content: {str(e)}")

    @classmethod
    def _extract_text_from_docx(cls, content: bytes) -> str:
        """Extracts text from DOCX bytes using python-docx."""
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n".join(full_text)
        except Exception as e:
            raise ResumeParserError(f"Corrupted or invalid DOCX document: {str(e)}")
